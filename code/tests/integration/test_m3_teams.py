"""M3: importers, research provenance, Planning/Requirements operations, diagrams, workflow.

Slices of AT-005 (epistemic labels), AT-006 (flexible documents), AT-007 (invalid
handoffs rejected), AT-009 (student adds and revises diagrams), AT-010 (manual edits
preserved with provenance) and the M3 exit evidence: PDF/Markdown brief -> traceable
plan -> SRS through the real services with a recorded fixture.
"""

from __future__ import annotations

import asyncio
import io
import json
from pathlib import Path

import pytest

from sdlc.adapters.brains.fixture import FixtureScript
from sdlc.adapters.registry import AdapterRegistry
from sdlc.application.handoffs import Actor, ActorKind, HandoffService
from sdlc.application.ingress import import_material
from sdlc.application.views import TeamView
from sdlc.application.workflow import WorkflowService
from sdlc.artifacts.diagrams import recommend, starter_source, validate_source
from sdlc.artifacts.imports import extract
from sdlc.cli.main import main
from sdlc.domain.errors import DeniedScope, InvalidSchema, MissingInput
from sdlc.domain.models import ArtifactRef, BundleKind, DiagramRequest, DiagramType, Origin, OwnerTeam, Task, TeamId
from sdlc.domain.states import TaskState
from sdlc.runtime.scheduler import Scheduler
from sdlc.verification.canonical import load_json_strict
from sdlc.verification.outputs import check_candidates
from sdlc.domain.models import ArtifactCandidate

ROOT = Path(__file__).resolve().parents[2]
EXAMPLE = ROOT / "examples" / "library-reservation"
PARTICIPANT = Actor(ActorKind.PARTICIPANT, "instructor")


# ----------------------------------------------------------------------------- helpers


def make_pdf(pages: list[str | None]) -> bytes:
    """Minimal valid PDF (with xref) where None pages have no text stream (like a scan)."""
    objects: list[bytes] = []
    font_id = 3 + 2 * len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for i, text in enumerate(pages):
        page_id, content_id = 3 + 2 * i, 4 + 2 * i
        if text is None:
            objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
            objects.append(b"<< /Length 0 >>\nstream\n\nendstream")
        else:
            stream = f"BT /F1 12 Tf 72 712 Td ({text}) Tj ET".encode()
            objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
            objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("요구사항 명세서", level=1)
    document.add_paragraph("회원은 도서를 예약할 수 있다. A member can reserve a copy.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "ID", "Statement"
    table.cell(1, 0).text, table.cell(1, 1).text = "REQ-9", "Reject a second active reservation"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------------------- extraction (AT-006)


def test_pdf_extraction_locates_pages_and_flags_scanned_pages():
    complete = extract(make_pdf(["Reserve one copy per member", "Cancel own reservation"]), "application/pdf")
    assert complete.completeness == "complete" and [l.index for l in complete.locators] == [1, 2]
    assert "Reserve one copy" in complete.locators[0].text
    partial = extract(make_pdf(["Text page", None]), "application/pdf")
    assert partial.completeness == "partial" and partial.units_empty == 1 and "page 2" in partial.warnings[0]
    scanned = extract(make_pdf([None]), "application/pdf")
    assert scanned.completeness == "unreadable"


def test_docx_extraction_keeps_headings_tables_and_korean():
    result = extract(make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    kinds = [l.kind for l in result.locators]
    assert kinds[0] == "heading" and "table_cell" in kinds and result.completeness == "complete"
    assert "요구사항 명세서" in result.text and "REQ-9" in result.text
    cell = next(l for l in result.locators if l.text == "REQ-9")
    assert cell.detail == "table 1, row 2, col 1"


def test_import_pdf_partial_yields_needs_input_not_empty_success(context, tmp_path):
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(make_pdf(["Readable page", None, None]))
    result = import_material(context, team=TeamId.PLANNING, path=pdf, purpose="scanned brief")
    assert result.receipt.status.value == "accepted"  # transport ok, original retained
    assert result.extraction == "partial" and result.needs_input and "2 of 3" in result.needs_input
    view = TeamView(context, TeamId.PLANNING)
    extraction = load_json_strict(view.read_text(result.extraction_artifact.ref, [result.delivery.delivery_id]))
    assert extraction["completeness"] == "partial" and extraction["source"] == result.artifact.ref.to_json_dict()
    assert [l["index"] for l in extraction["locators"]] == [1]


def test_import_docx_roundtrips_unicode(context, tmp_path):
    path = tmp_path / "spec.docx"
    path.write_bytes(make_docx())
    result = import_material(context, team=TeamId.REQUIREMENTS, path=path, purpose="spec")
    assert result.extraction == "complete" and result.needs_input is None
    text = TeamView(context, TeamId.REQUIREMENTS).read_text(result.extraction_artifact.ref, [result.delivery.delivery_id])
    assert "요구사항 명세서" in text


# ----------------------------------------------------------------------------- output checks (AT-007, AT-005)


def _candidate(kind: str, text: str, media="text/markdown", path="docs/x.md") -> ArtifactCandidate:
    return ArtifactCandidate(kind=kind, media_type=media, relative_path=path, content_text=text, workspace_path=None)


def test_output_checks_reject_fabricated_evidence_and_invalid_records():
    fake_run = _candidate("execution_record", "{}", "application/json", "runs/1.json")
    assert [f.code for f in check_candidates([fake_run])] == ["fabricated_evidence"]
    bad_req = _candidate("requirements", json.dumps({"schema_version": "1.0", "scope": "x", "requirements": [], "claims": [], "assumptions": []}), "application/json", "requirements.json")
    assert [f.code for f in check_candidates([bad_req])] == ["invalid_record"]
    plan = _candidate("project_plan", "# Plan\n\n## Problem\nx\n", path="PROJECT_PLAN.md")
    assert [f.code for f in check_candidates([plan])] == ["missing_sections"]
    good = _candidate("diagram_source", "flowchart TD\n  A --> B\n", "text/vnd.mermaid", "d.mmd")
    assert check_candidates([good]) == []
    bad = _candidate("diagram_source", "<svg/>", "image/svg+xml", "d.svg")
    assert [f.code for f in check_candidates([bad])] == ["unsupported_diagram_syntax"]


# ----------------------------------------------------------------------------- end-to-end: brief -> plan -> SRS


@pytest.fixture
def walkthrough(context):
    script = load_json_strict((EXAMPLE / "fixture-script.json").read_text(encoding="utf-8"))
    adapters = AdapterRegistry(context.config, context.root, fixture_script=FixtureScript.from_raw(script), fixture_research=script["research"])
    scheduler = Scheduler(context, adapters, owner_id="coord")
    return {"ctx": context, "scheduler": scheduler, "workflow": WorkflowService(context, scheduler), "handoffs": HandoffService(context), "script": script}


def _fill_inbox(ctx, script: dict, team: str, delivery_id: str) -> None:
    """The recorded task graphs use the placeholder @INBOX@ for the delivery id."""
    for key in (f"operation:{team}.plan_work",):
        step = script["steps"][key]
        for candidate in step["proposal"]["artifact_candidates"]:
            candidate["content_text"] = candidate["content_text"].replace("@INBOX@", delivery_id)


def test_brief_to_plan_to_srs_with_provenance(walkthrough, tmp_path):
    ctx, scheduler, workflow, handoffs, script = (walkthrough[k] for k in ("ctx", "scheduler", "workflow", "handoffs", "script"))
    # 1. participant imports the brief for planning
    imported = import_material(ctx, team=TeamId.PLANNING, path=EXAMPLE / "brief.md", purpose="course brief")
    _fill_inbox(ctx, script, "planning", imported.delivery.delivery_id)
    scheduler.adapters._brains.clear()  # rebuild the fixture brain with the filled script
    scheduler.adapters._fixture_script = FixtureScript.from_raw(script)
    # 2. leader proposes, participant approves, runtime runs
    plan = workflow.propose_stage_plan(TeamId.PLANNING, objective="Plan the library reservation demo")
    [outcome] = asyncio.run(scheduler.run_ready(TeamId.PLANNING))
    assert outcome.task_id == plan.task_id and outcome.state == TaskState.COMPLETED
    with pytest.raises(MissingInput):
        workflow.approve_plan("nope", PARTICIPANT)
    approved = workflow.approve_plan(plan.task_id, PARTICIPANT)
    assert approved["tasks"] == ["planning-frame", "planning-research", "planning-handoff"]
    outcomes = {o.task_id: o for o in asyncio.run(scheduler.run_ready(TeamId.PLANNING))}
    assert {o.state for o in outcomes.values()} == {TaskState.COMPLETED}
    # research evidence was registered by the runtime with fixture origin; summary cites it
    evidence = [a for a in scheduler.artifacts.list(OwnerTeam.PLANNING) if a.kind == "research_evidence"]
    assert evidence and evidence[0].origin.value == "fixture" and evidence[0].created_by_task == "planning-research"
    summary = next(a for a in scheduler.artifacts.list(OwnerTeam.PLANNING) if a.kind == "research_summary")
    assert "synthetic" in scheduler.artifacts.read_bytes(summary.ref).decode().lower()
    # provisional: planning has no accepted baseline yet, so outputs are draft with a provisional finding
    assert any(f.code == "provisional" for f in outcomes["planning-frame"].result.findings)
    assert scheduler.artifacts.status(outcomes["planning-frame"].artifacts[0]) == "draft"
    # 3. planning publishes its handoff: plan + original brief (forwarded) -> requirements
    plan_doc = next(a for a in scheduler.artifacts.list(OwnerTeam.PLANNING) if a.kind == "project_plan")
    with pytest.raises(MissingInput, match="dependency closure"):
        handoffs.publish(source=OwnerTeam.PLANNING, target=TeamId.REQUIREMENTS, kind=BundleKind.FORWARD, bundle_id="b-plan", cycle_id="cycle-001",
                         baseline_id="baseline-planning-001", artifact_refs=[plan_doc.ref], summary="plan without its inputs")
    manifest = [imported.artifact.ref, imported.extraction_artifact.ref, plan_doc.ref, summary.ref] + [e.ref for e in evidence]
    _, bundle_ref = handoffs.publish(source=OwnerTeam.PLANNING, target=TeamId.REQUIREMENTS, kind=BundleKind.FORWARD, bundle_id="b-plan", cycle_id="cycle-001",
                                     baseline_id="baseline-planning-001", artifact_refs=manifest, summary="Project plan and the original brief")
    delivery = handoffs.deliver(bundle_ref=bundle_ref, expected_target_baseline=None)
    assert handoffs.receive(delivery.delivery_id).status.value == "accepted"
    handoffs.accept_baseline(team=TeamId.REQUIREMENTS, delivery_id=delivery.delivery_id, actor=PARTICIPANT, rationale="plan reviewed in class")
    # 4. requirements team: propose -> approve -> run in ACCEPTED mode
    _fill_inbox(ctx, script, "requirements", delivery.delivery_id)
    scheduler.adapters._brains.clear()
    scheduler.adapters._fixture_script = FixtureScript.from_raw(script)
    plan2 = workflow.propose_stage_plan(TeamId.REQUIREMENTS, objective="Specify requirements")
    asyncio.run(scheduler.run_ready(TeamId.REQUIREMENTS))
    workflow.approve_plan(plan2.task_id, PARTICIPANT)
    outcomes = {o.task_id: o for o in asyncio.run(scheduler.run_ready(TeamId.REQUIREMENTS))}
    assert {o.state for o in outcomes.values()} == {TaskState.COMPLETED}, {k: v.reason for k, v in outcomes.items()}
    assert not any(f.code == "provisional" for f in outcomes["req-extract"].result.findings)
    req_set = next(a for a in scheduler.artifacts.list(OwnerTeam.REQUIREMENTS) if a.kind == "requirements")
    assert scheduler.artifacts.status(req_set.ref) == "reviewable"
    parsed = load_json_strict(scheduler.artifacts.read_bytes(req_set.ref).decode())
    assert [r["requirement_id"] for r in parsed["requirements"]] == ["REQ-001", "REQ-002", "REQ-003"]
    assert all(c["source_locators"] for c in parsed["claims"])  # traceable to the brief
    assert req_set.depends_on  # provenance back to granted inputs
    # Requirements cannot read planning's private problem frame (not in the manifest)
    frame = next(a for a in scheduler.artifacts.list(OwnerTeam.PLANNING) if a.kind == "problem_frame")
    with pytest.raises(DeniedScope):
        TeamView(ctx, TeamId.REQUIREMENTS).read_text(frame.ref, [delivery.delivery_id])


def test_requirements_policy_rejects_synthetic_only_must(walkthrough):
    from sdlc.domain.models import Task, TaskLimits, WorkerProposal
    from sdlc.teams.requirements.operations import RequirementsPolicy

    policy = RequirementsPolicy(name="requirements.extract.v1", short_name="extract")
    req_set = {
        "schema_version": "1.0", "scope": "x", "assumptions": [],
        "claims": [{"claim_id": "c1", "statement": "persona wants it", "basis": "synthetic", "source_refs": [], "source_locators": [], "limitations": []}],
        "requirements": [{"requirement_id": "REQ-1", "revision": 1, "kind": "functional", "statement": "s", "rationale": "r", "priority": "must",
                          "source_claim_ids": ["c1"], "acceptance": ["observable outcome happens"], "verification_method": "inspection", "open_questions": []}],
    }
    proposal = WorkerProposal(schema_version="1.0", task_id="t", summary="s", artifact_candidates=[_candidate("requirements", json.dumps(req_set), "application/json", "requirements.json")], findings=[], questions=[], suggested_next_operations=[])
    task = Task(schema_version="1.0", project_id="p", cycle_id="c", task_id="t", team_id=TeamId.REQUIREMENTS, agent_instance_id="a", operation="requirements.extract", operation_version=1,
                input_deliveries=["d"], input_refs=[], depends_on=[], parameters={}, expected_output_schema="WorkerProposal", backend_profile="fixture", workspace_ref=None,
                limits=TaskLimits(max_attempts=1, timeout_seconds=1, max_output_tokens=1, max_cost_microunits=None))
    codes = [f.code for f in policy.check_outputs(task, proposal)]
    assert "synthetic_only_basis" in codes


# ----------------------------------------------------------------------------- diagrams (AT-009, AT-010)


def test_recommendations_are_selective_and_explained():
    text = (EXAMPLE / "brief.md").read_text(encoding="utf-8")
    recs = recommend(TeamId.REQUIREMENTS, text)
    types = [r.diagram_type for r in recs]
    assert DiagramType.USE_CASE in types and DiagramType.SEQUENCE in types and len(recs) <= 5
    assert all(r.question and r.reason for r in recs)
    assert recommend(TeamId.PLANNING, "no dates here")[0].diagram_type == DiagramType.CONTEXT


def test_starter_sources_are_valid():
    for kind in DiagramType:
        media, source = starter_source(kind, ["REQ-001"], "test")
        assert validate_source(source, media) == [], kind


def test_diagram_request_run_status_and_participant_revision(walkthrough, tmp_path):
    ctx, scheduler, workflow = walkthrough["ctx"], walkthrough["scheduler"], walkthrough["workflow"]
    imported = import_material(ctx, team=TeamId.REQUIREMENTS, path=EXAMPLE / "brief.md", purpose="brief")
    request = DiagramRequest(schema_version="1.0", project_id=ctx.project_id, team_id=TeamId.REQUIREMENTS, requested_by="participant",
                             source_refs=[imported.artifact.ref], diagram_type=DiagramType.USE_CASE, purpose="actor goals", scope_ids=["REQ-001"], preferred_renderer="plantuml")
    with pytest.raises(DeniedScope):
        workflow.request_diagram(request, input_deliveries=[])  # source not granted without the delivery
    task_id = workflow.request_diagram(request, input_deliveries=[imported.delivery.delivery_id])
    [outcome] = asyncio.run(scheduler.run_ready(TeamId.REQUIREMENTS))
    assert outcome.task_id == task_id and outcome.state == TaskState.COMPLETED, outcome.reason
    [ref] = outcome.artifacts
    status = workflow.diagram_status(ref)
    assert status["syntax_errors"] == [] and status["rendered"] is False and status["label"] == "unrendered"
    # participant edits the source: new revision, participant origin, original superseded, nothing overwritten
    edited = scheduler.artifacts.read_bytes(ref).decode().replace('usecase "Cancel own reservation" as UC2', 'usecase "Cancel own reservation (REQ-004?)" as UC2').encode()
    revised = workflow.revise_artifact(ref, edited, actor_id="student-1", note="added question about cancellation")
    assert revised.revision == 2 and revised.origin.value == "participant" and revised.depends_on == [ref]
    assert scheduler.artifacts.read_bytes(ref) != edited and scheduler.artifacts.status(ref) == "superseded"
    with pytest.raises(InvalidSchema):
        workflow.revise_artifact(revised.ref, b"not a diagram", actor_id="student-1", note="broken")
    # an unsupported notation request produces a capability finding, not a lookalike
    bad = ArtifactCandidate(kind="diagram_source", media_type="image/png", relative_path="d.png", content_text="...", workspace_path=None)
    assert [f.code for f in check_candidates([bad])] == ["unsupported_diagram_syntax"]


# ----------------------------------------------------------------------------- CLI


def run(capsys, *argv) -> tuple[int, dict]:
    code = main(["--json", *argv])
    return code, json.loads(capsys.readouterr().out)


def test_cli_walkthrough_planning_stage(tmp_path, capsys):
    home = str(tmp_path / "home")
    run(capsys, "init", "demo", "--home", home)
    code, env = run(capsys, "import", "--project", "demo", "--home", home, "--team", "planning", "--file", str(EXAMPLE / "brief.md"))
    assert code == 0 and env["data"]["extraction"] == "complete" and env["data"]["needs_input"] is None
    delivery_id = env["data"]["delivery_id"]
    script = json.loads((EXAMPLE / "fixture-script.json").read_text(encoding="utf-8"))
    for candidate in script["steps"]["operation:planning.plan_work"]["proposal"]["artifact_candidates"]:
        candidate["content_text"] = candidate["content_text"].replace("@INBOX@", delivery_id)
    script_path = tmp_path / "script.json"
    script_path.write_text(json.dumps(script), encoding="utf-8")

    code, env = run(capsys, "stage", "propose", "--project", "demo", "--home", home, "--team", "planning", "--objective", "plan it", "--script", str(script_path))
    assert code == 0, env
    plan_task = env["data"]["plan_task_id"]
    code, env = run(capsys, "stage", "run", "--project", "demo", "--home", home, "--script", str(script_path))
    assert code == 0 and env["data"]["outcomes"][0]["state"] == "completed", env
    code, env = run(capsys, "stage", "approve", "--project", "demo", "--home", home, "--plan-task", plan_task)
    assert code == 0 and len(env["data"]["tasks"]) == 3, env
    code, env = run(capsys, "stage", "run", "--project", "demo", "--home", home, "--script", str(script_path))
    assert code == 0 and {o["state"] for o in env["data"]["outcomes"]} == {"completed"}, env
    code, env = run(capsys, "diagram", "recommend", "--project", "demo", "--home", home, "--team", "planning", "--source", env["data"]["outcomes"][0]["artifacts"][0] if False else _import_ref(capsys, home), "--delivery", delivery_id)
    assert code == 0 and env["data"]["recommendations"][0]["type"] == "context", env
    code, env = run(capsys, "artifacts", "list", "--project", "demo", "--home", home, "--owner", "planning")
    kinds = {a["kind"] for a in env["data"]["artifacts"]}
    assert {"task_graph", "problem_frame", "research_evidence", "research_summary", "project_plan"} <= kinds


def _import_ref(capsys, home) -> str:
    code, env = run(capsys, "artifacts", "list", "--project", "demo", "--home", home, "--owner", "participant")
    return next(a["ref"] for a in env["data"]["artifacts"] if a["kind"] == "import_original")


def test_design_gate_rejects_invented_requirement_ids(walkthrough):
    """FR-008: a DesignSet must pin the delivered RequirementsSet and cite only its ids."""
    from sdlc.domain.models import TaskLimits

    ctx, scheduler, handoffs = walkthrough["ctx"], walkthrough["scheduler"], walkthrough["handoffs"]
    # A requirements artifact (RequirementsSet JSON) delivered to design and accepted as baseline.
    req_json = json.loads((EXAMPLE / "fixture-script.json").read_text(encoding="utf-8"))["steps"]["operation:requirements.extract"]["proposal"]["artifact_candidates"][0]["content_text"]
    req = scheduler.artifacts.register(owner=OwnerTeam.REQUIREMENTS, artifact_id="req-set", kind="requirements", media_type="application/json",
                                       relative_path="requirements.json", data=req_json.encode("utf-8"), origin=Origin.FIXTURE)
    _, ref = handoffs.publish(source=OwnerTeam.REQUIREMENTS, target=TeamId.DESIGN, kind=BundleKind.FORWARD, bundle_id="b-req", cycle_id="cycle-001",
                              baseline_id="baseline-req", artifact_refs=[req.ref], summary="requirements")
    delivery = handoffs.deliver(bundle_ref=ref, expected_target_baseline=None)
    handoffs.receive(delivery.delivery_id)
    handoffs.accept_baseline(team=TeamId.DESIGN, delivery_id=delivery.delivery_id, actor=PARTICIPANT, rationale="ok")

    def design_set(requirements_ref, ids):
        return json.dumps({"schema_version": "1.0", "requirements_ref": requirements_ref, "decisions": [], "open_questions": [],
                           "elements": [{"element_id": "CMP-1", "kind": "component", "name": "Reservation service", "description": "d", "requirement_ids": ids, "contract_refs": [], "checks": []}]})

    arch = {"kind": "architecture", "media_type": "text/markdown", "relative_path": "docs/ARCHITECTURE.md", "workspace_path": None,
            "content_text": "# Architecture\n## Interface\n## Behavior\n## Requirement mapping\n"}
    bad = {"kind": "design", "media_type": "application/json", "relative_path": "design.json", "workspace_path": None,
           "content_text": design_set({"artifact_id": "req-set", "revision": 1, "sha256": req.sha256}, ["REQ-001", "REQ-X-9"])}
    # right id and revision but a mangled hash: the runtime normalises it to the pinned value
    good = dict(bad, content_text=design_set({"artifact_id": "req-set", "revision": 1, "sha256": "f" * 64}, ["REQ-001", "REQ-002", "REQ-003"]))
    script = {"steps": {"task:d-arch": [{"proposal": {"summary": "bad ids", "artifact_candidates": [bad, arch]}},
                                        {"proposal": {"summary": "fixed", "artifact_candidates": [good, arch]}}]}}
    scheduler.adapters._brains.clear()
    scheduler.adapters._fixture_script = FixtureScript.from_raw(script)
    task = Task(schema_version="1.0", project_id=ctx.project_id, cycle_id="cycle-001", task_id="d-arch", team_id=TeamId.DESIGN, agent_instance_id="design-architect",
                operation="design.define_architecture", operation_version=1, input_deliveries=[delivery.delivery_id], input_refs=[req.ref], depends_on=[],
                parameters={"objective": "architecture", "scope_ids": [], "requested_artifact_kinds": ["design"], "diagram_request": None},
                expected_output_schema="#/$defs/WorkerProposal", backend_profile="fixture", workspace_ref=None,
                limits=TaskLimits(max_attempts=3, timeout_seconds=10, max_output_tokens=100, max_cost_microunits=0))
    scheduler.graphs.admit(TeamId.DESIGN, [task])
    first = asyncio.run(scheduler.run_ready(TeamId.DESIGN))[0]
    assert first.state == TaskState.RETRY_WAIT and "unknown_requirement_id" in first.reason and "REQ-X-9" in first.reason
    second = asyncio.run(scheduler.run_ready(TeamId.DESIGN))[0]
    assert second.state == TaskState.COMPLETED, second.reason
    assert scheduler.artifacts.status(second.artifacts[0]) == "reviewable"  # accepted mode: inputs are the baseline
    stored = json.loads(scheduler.artifacts.read_bytes(next(r for r in second.artifacts if r.artifact_id.endswith("design"))).decode())
    assert stored["requirements_ref"]["sha256"] == req.sha256
    assert any(f.code == "ref_normalized" for f in second.result.findings)


def test_handoff_document_cannot_renumber_requirements(walkthrough):
    """A REQUIREMENTS.md citing ids absent from the canonical RequirementsSet is rejected."""
    from sdlc.domain.models import WorkerProposal
    from sdlc.teams.common import CommonLifecyclePolicy

    script = walkthrough["script"]
    proposal_raw = script["steps"]["operation:requirements.prepare_handoff"]["proposal"]
    candidates = [ArtifactCandidate.model_validate(c) for c in proposal_raw["artifact_candidates"]]
    policy = CommonLifecyclePolicy(name="requirements.prepare_handoff.v1", short_name="prepare_handoff", team="requirements")
    task = Task(schema_version="1.0", project_id="p", cycle_id="c", task_id="t", team_id=TeamId.REQUIREMENTS, agent_instance_id="a", operation="requirements.prepare_handoff",
                operation_version=1, input_deliveries=["d"], input_refs=[], depends_on=[], parameters={}, expected_output_schema="WorkerProposal", backend_profile="fixture",
                workspace_ref=None, limits=__import__("sdlc.domain.models", fromlist=["TaskLimits"]).TaskLimits(max_attempts=1, timeout_seconds=1, max_output_tokens=1, max_cost_microunits=None))
    good = WorkerProposal(schema_version="1.0", task_id="t", summary="s", artifact_candidates=candidates, findings=[], questions=[], suggested_next_operations=[])
    assert [f.code for f in policy.check_outputs(task, good)] == []
    renumbered = [c.model_copy(update={"content_text": c.content_text.replace("REQ-001", "REQ-F-01")}) if c.kind == "requirements_document" else c for c in candidates]
    bad = good.model_copy(update={"artifact_candidates": renumbered})
    codes = [f.code for f in policy.check_outputs(task, bad)]
    assert "unknown_requirement_id" in codes
    # The JSON record is forwarded by the runtime's manifest; the document alone is fine.
    without_json = good.model_copy(update={"artifact_candidates": [c for c in candidates if c.kind != "requirements"]})
    assert [f.code for f in policy.check_outputs(task, without_json)] == []


def test_cycle_run_automates_planning_to_requirements_and_stops_when_a_person_is_needed(walkthrough):
    """Continuous mode: policy decisions drive propose/approve/run/handoff; stops on needs_input."""
    from sdlc.application.cycle import CycleRunner

    ctx, scheduler = walkthrough["ctx"], walkthrough["scheduler"]
    ctx.config = ctx.config.model_copy(update={"project": ctx.config.project.model_copy(update={"mode": "continuous"})})
    import_material(ctx, team=TeamId.PLANNING, path=EXAMPLE / "brief.md", purpose="course brief")
    runner = CycleRunner(ctx, scheduler)
    report = runner.run(start=TeamId.PLANNING, end=TeamId.REQUIREMENTS)
    assert not report.stopped, report.to_dict()
    assert [s.status for s in report.stages] == ["completed", "completed"]
    assert HandoffService(ctx).baseline(TeamId.REQUIREMENTS)[0] is not None
    req = [a for a in scheduler.artifacts.list(OwnerTeam.REQUIREMENTS) if a.kind == "requirements"]
    assert req and scheduler.artifacts.status(req[0].ref) == "reviewable"
    decisions = [e for e in repo_events(ctx, "decision.recorded")]
    assert decisions and all(e["payload"]["actor_kind"] == "policy" for e in decisions)
    # Re-running is idempotent for finished stages and continues into design, where the
    # fixture has no recorded leader reply: the runner stops and says what to do.
    report = runner.run(start=TeamId.PLANNING, end=TeamId.DESIGN)
    assert report.stopped and "design" in (report.next_action or "")
    assert [s.status for s in report.stages][:2] == ["completed", "completed"]


def repo_events(ctx, event_type):
    from sdlc.storage import repositories as repo

    return repo.list_events(ctx.db, ctx.project_id, event_type)


def test_cycle_run_refuses_collaborative_mode(walkthrough):
    from sdlc.application.cycle import CycleRunner
    from sdlc.domain.errors import Conflict

    with pytest.raises(Conflict, match="--ask"):
        CycleRunner(walkthrough["ctx"], walkthrough["scheduler"])  # collaborative mode without a gate callback


def test_everyday_commands_new_run_show_export(tmp_path, capsys):
    """sdlc new / run --until / show / export / brain / continue on the fixture brain."""
    home = str(tmp_path / "home")
    script = json.loads((EXAMPLE / "fixture-script.json").read_text(encoding="utf-8"))
    script_path = tmp_path / "script.json"
    script_path.write_text(json.dumps(script), encoding="utf-8")

    code, env = run(capsys, "new", "demo", "--home", home, "--brief", str(EXAMPLE / "brief.md"), "--brain", "fixture")
    assert code == 0 and env["data"]["mode"] == "auto" and env["data"]["next"] == "sdlc run demo", env

    code, env = run(capsys, "show", "demo", "--home", home)
    assert code == 0 and env["data"]["stages"][0]["state"] == "ready to start" and "sdlc run demo" in env["data"]["what_to_do"]

    code, env = run(capsys, "run", "demo", "--home", home, "--until", "requirements", "--script", str(script_path))
    assert code == 0 and not env["data"]["stopped"], env
    assert [s["stage"] for s in env["data"]["stages"]] == ["planning", "requirements"]

    code, env = run(capsys, "show", "demo", "--home", home)
    states = {s["stage"].split(" ")[0].lower(): s["state"] for s in env["data"]["stages"]}
    assert states["planning"] == "finished" and states["requirements"] == "finished" and states["design"] == "waiting for input"
    assert "handed on" in env["data"]["what_to_do"]  # --until leaves the package for you to review first

    code, env = run(capsys, "export", "demo", "--home", home, "--to", str(tmp_path / "out"))
    assert code == 0 and (tmp_path / "out" / "requirements" / "REQUIREMENTS.md").exists() and (tmp_path / "out" / "participant" / "imports").exists()

    # the next run hands the requirements package to design first; design has no recorded
    # leader reply in the fixture, so the run then stops with a readable reason
    code, env = run(capsys, "run", "demo", "--home", home, "--step", "--script", str(script_path))
    assert code == 0 and env["data"]["stopped"] and "design" in env["data"]["what_to_do"], env
    assert env["data"]["stages"][0]["stage"] == "requirements->design" and env["data"]["stages"][0]["status"] == "completed"
    code, env = run(capsys, "show", "demo", "--home", home)
    assert "blocked" in env["data"]["what_to_do"]

    code, env = run(capsys, "brain", "demo", "fixture", "--home", home, "--stage", "design")
    assert code == 0 and env["data"]["brain"]["design"] == "fixture"

    # ask mode: a fresh project whose first gate needs a terminal; without one it refuses clearly
    code, env = run(capsys, "new", "demo2", "--home", home, "--brief", str(EXAMPLE / "brief.md"), "--brain", "fixture", "--mode", "ask")
    assert code == 0 and env["data"]["next"] == "sdlc run demo2 --ask"
    code, env = run(capsys, "run", "demo2", "--home", home, "--script", str(script_path))
    assert code == 1 and "--ask" in env["errors"][0]["message"]  # ask-mode project refuses an unattended run
    code, env = run(capsys, "run", "demo2", "--home", home, "--ask", "--script", str(script_path))
    assert code == 1 and "interactive" in env["errors"][0]["message"]
